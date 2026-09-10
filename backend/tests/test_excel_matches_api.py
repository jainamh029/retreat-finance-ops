"""Regression guard: retreat_finance_model.xlsx must not silently drift from the API.

The workbook's aggregate cells (aging buckets, DSO/DPO, cash-flow ending balances) are *live
formulas*. This test builds the workbook fresh, has LibreOffice headless **recalculate** it, then
reads the recalculated values with openpyxl `data_only=True` and compares them to the same
`backend.data_access` calls the FastAPI layer serves. If cashflow.py / reconcile.py / aging.py
ever change in a way the Excel formulas don't mirror, this fails immediately.

Skipped (not failed) when LibreOffice or the built DB is unavailable, so the pure unit suite
still runs on a machine without either.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from backend import data_access as dao

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import build_excel_model as bem  # noqa: E402

_SOFFICE_CANDIDATES = [
    "soffice", "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice", "/usr/bin/libreoffice",
]


def _find_soffice() -> str | None:
    for c in _SOFFICE_CANDIDATES:
        p = shutil.which(c) or (c if os.path.exists(c) else None)
        if p:
            return p
    return None


@pytest.fixture(scope="module")
def recalced_wb(real_db, tmp_path_factory):
    soffice = _find_soffice()
    if soffice is None:
        pytest.skip("LibreOffice (soffice) not found — cannot recalculate the workbook")
    work = tmp_path_factory.mktemp("xlsx")
    src = bem.build(work / "model.xlsx")
    out = work / "out"
    out.mkdir()
    env = {**os.environ, "HOME": str(work)}  # isolated LO profile
    res = subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(out), str(src)],
        capture_output=True, text=True, timeout=240, env=env,
    )
    recalced = out / "model.xlsx"
    if res.returncode != 0 or not recalced.exists():
        pytest.skip(f"LibreOffice recalc failed: {res.stderr.strip()[:200]}")
    return openpyxl.load_workbook(recalced, data_only=True)


def _row(ws, label, col=2):
    for r in ws.iter_rows():
        if r and str(r[0].value).strip() == label:
            return r[col - 1].value
    raise AssertionError(f"row {label!r} not found in {ws.title}")


# --------------------------------------------------------------------------------------
def test_dashboard_dso_dpo_recalc_matches_api(recalced_wb):
    api_ar, api_ap = dao.aging_ar(), dao.aging_ap()
    assert _row(recalced_wb["AR"], "DSO = open / trailing × 365") == pytest.approx(
        api_ar["metric"]["value"], abs=0.15)
    assert _row(recalced_wb["AP"], "DPO = open / trailing × 365") == pytest.approx(
        api_ap["metric"]["value"], abs=0.15)
    # and the Dashboard tab's own cells
    assert recalced_wb["Dashboard"]["B7"].value == pytest.approx(api_ar["metric"]["value"], abs=0.15)
    assert recalced_wb["Dashboard"]["B8"].value == pytest.approx(api_ap["metric"]["value"], abs=0.15)


def test_aging_buckets_recalc_match_api(recalced_wb):
    for kind, pay in (("AR", dao.aging_ar()), ("AP", dao.aging_ap())):
        for b in ("current", "0-30", "31-60", "61-90", "90+"):
            api_amt = next(x["amount"] for x in pay["buckets"] if x["bucket"] == b)
            assert _row(recalced_wb[kind], b) == pytest.approx(api_amt, abs=0.02), f"{kind} {b}"


def test_cash_position_and_match_rate(recalced_wb):
    summ = dao.dashboard_summary()
    assert recalced_wb["Dashboard"]["B9"].value == pytest.approx(summ["cards"]["cash_position"], abs=0.02)
    assert recalced_wb["Dashboard"]["B12"].value == pytest.approx(
        summ["reconciliation"]["match_rate_bank"] / 100, abs=0.001)


def test_cashflow_13_week_endings_recalc_match_api(recalced_wb):
    fc = dao.cashflow_forecast()
    cf = recalced_wb["Cash Flow Forecast"]
    for i, wk in enumerate(fc["weeks"]):
        assert cf.cell(row=6 + i, column=7).value == pytest.approx(wk["ending_balance"], abs=0.02), f"wk{i+1}"
        assert cf.cell(row=6 + i, column=10).value == "ok", f"wk{i+1} check column"
    # wk-13 also equals the dashboard's projected cash position
    assert cf.cell(row=18, column=7).value == pytest.approx(
        dao.dashboard_summary()["cards"]["cash_position_week13"], abs=0.02)


def test_budget_variance_recalc_matches_api(recalced_wb):
    rlist = dao.retreats_view()["rows"]
    spot = max((r for r in rlist if r["has_actuals"]), key=lambda r: abs(r["variance_abs"]))
    api = dao.budget_vs_actual(spot["retreat_id"])["total"]
    ws = recalced_wb["Budget vs Actual"]
    rows = list(ws.iter_rows())
    total_var = None
    for i, r in enumerate(rows):
        if r and str(r[0].value or "").startswith(spot["retreat_id"]):
            for j in range(i, min(i + 9, len(rows))):
                if rows[j][0].value == "TOTAL":
                    total_var = rows[j][3].value          # column D = variance $
    assert total_var == pytest.approx(round(api["actual"] - api["budget"], 2), abs=0.02)
    assert total_var == pytest.approx(spot["variance_abs"], abs=0.02)
